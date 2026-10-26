"""Export utilities: Excel (.xlsx) and PDF report generation."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Column display configs per phase ──────────────────────────────────────

# Index (#) is UI-only for user feedback; excluded from export.
PHASE_1_COLS = {
    "input_risk_mapping_id": "UI ID",
    "source_document": "Source",
    "source_ref": "Ref",
    "requirement_theme": "Theme",
    "risk_id": "Risk ID",
    "risk_statement": "Risk Statement",
    "risk_category": "Category",
    "severity": "Severity",
    "why_identified": "Rationale",
    "source_quote": "Quote",
    "user_status": "Status",
    "Comment": "Comment",
}

PHASE_2_COLS = {
    "input_risk_mapping_id": "UI ID",
    "risk_id": "Risk ID",
    "risk_statement": "Risk Statement",
    "control_id": "Control ID",
    "control_description": "Control Description",
    "why_control_assigned": "Rationale",
    "control_type": "Type",
    "execution_type": "Execution",
    "user_status": "Status",
    "Comment": "Comment",
}

PHASE_3_COLS = {
    "control_id": "Control ID",
    "control_description": "Control Description",
    "process_id": "Process ID",
    "process_name": "Process Step",
    "process_step": "Process Detail",
    "why_process_assigned": "Rationale",
    "user_status": "Status",
    "Comment": "Comment",
}

SEVERITY_FILLS = {
    "High": PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid"),
    "Medium": PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid"),
    "Low": PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"),
}

HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
CELL_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)

# ── PDF column width weights ──────────────────────────────────────────────
# Higher weight = wider column.  Unspecified keys default to 2.
_COL_WEIGHT: dict[str, float] = {
    "input_risk_mapping_id": 1.0,
    "risk_control_mapping_id": 1.0,
    "control_process_mapping_id": 1.0,
    "risk_id": 1.0,
    "control_id": 1.2,
    "process_id": 1.2,
    "source_ref": 1.2,
    "severity": 1.0,
    "control_type": 1.2,
    "execution_type": 1.2,
    "user_status": 1.0,
    "risk_category": 1.0,
    "source_document": 2.2,
    "requirement_theme": 2.0,
    "process_name": 2.5,
    "process_step": 2.5,
    "risk_statement": 4.0,
    "control_description": 4.0,
    "why_identified": 3.0,
    "why_control_assigned": 3.0,
    "why_process_assigned": 3.0,
    "source_quote": 4.0,
    "reason": 3.0,
    "status": 1.2,
    "Comment": 2.5,
}

_DEFAULT_WEIGHT = 2.0


def _cell_value(row: dict, key: str) -> str:
    """Get display value for a column; Comment falls back to feedback/user_notes."""
    if key == "Comment":
        return str(row.get("Comment") or row.get("feedback") or row.get("user_notes") or "")
    return str(row.get(key, ""))


# ═══════════════════════════════════════════════════════════════════════════
# Excel
# ═══════════════════════════════════════════════════════════════════════════

def _write_sheet(wb, sheet_name: str, data: list[dict], col_map: dict):
    """Write a formatted sheet into the workbook."""
    ws = wb.create_sheet(title=sheet_name)

    cols = [k for k in col_map if k == "Comment" or any(k in row for row in data)]
    headers = [col_map[k] for k in cols]

    for c_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    for r_idx, row in enumerate(data, 2):
        for c_idx, key in enumerate(cols, 1):
            value = _cell_value(row, key)
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.alignment = CELL_ALIGNMENT
            cell.border = THIN_BORDER

            if key == "severity" and value in SEVERITY_FILLS:
                cell.fill = SEVERITY_FILLS[value]

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    for c_idx, key in enumerate(cols, 1):
        max_len = len(headers[c_idx - 1])
        for row in data:
            val = _cell_value(row, key)
            max_len = max(max_len, min(len(val), 60))
        ws.column_dimensions[get_column_letter(c_idx)].width = max_len + 4


def build_excel(
    p1: list[dict] | None = None,
    p2: list[dict] | None = None,
    p3: list[dict] | None = None,
    verdict: dict | None = None,
    session_name: str = "",
) -> bytes:
    """Build a multi-sheet Excel workbook and return as bytes.
    Uses only approved phase data and verdict; feedback history and agent_message are not included.
    """
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    if p1:
        _write_sheet(wb, "Risk Register", p1, PHASE_1_COLS)
    if p2:
        _write_sheet(wb, "Risk-Control Matrix", p2, PHASE_2_COLS)
    if p3:
        _write_sheet(wb, "Governance Map", p3, PHASE_3_COLS)

    if verdict:
        ws = wb.create_sheet(title="Verdict")
        ws.cell(row=1, column=1, value="Overall Status").font = Font(bold=True)
        ws.cell(row=1, column=2, value=verdict.get("overall_status", ""))
        ws.cell(row=2, column=1, value="Verdict").font = Font(bold=True)
        ws.cell(row=2, column=2, value=verdict.get("overall_verdict", ""))
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 80

        risks = verdict.get("risks", [])
        if risks:
            ws.cell(row=4, column=1, value="Risk ID").font = HEADER_FONT
            ws.cell(row=4, column=1).fill = HEADER_FILL
            ws.cell(row=4, column=2, value="Status").font = HEADER_FONT
            ws.cell(row=4, column=2).fill = HEADER_FILL
            ws.cell(row=4, column=3, value="Reason").font = HEADER_FONT
            ws.cell(row=4, column=3).fill = HEADER_FILL
            ws.column_dimensions["C"].width = 60

            for i, r in enumerate(risks, 5):
                ws.cell(row=i, column=1, value=r.get("risk_id", ""))
                ws.cell(row=i, column=2, value=r.get("status", ""))
                ws.cell(row=i, column=3, value=r.get("reason", ""))

        recs = verdict.get("recommendations", [])
        if recs:
            start = 5 + len(risks) + 1
            ws.cell(row=start, column=1, value="Recommendations").font = Font(bold=True)
            for i, rec in enumerate(recs, start + 1):
                ws.cell(row=i, column=1, value=f"{i - start}.")
                ws.cell(row=i, column=2, value=rec)

    if not wb.sheetnames:
        wb.create_sheet(title="Empty")
        wb["Empty"].cell(row=1, column=1, value="No data available")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# PDF
# ═══════════════════════════════════════════════════════════════════════════

def _find_unicode_font() -> str | None:
    """Return path to a system TrueType font that supports Unicode."""
    from pathlib import Path
    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/System/Library/Fonts/Helvetica.ttc"),
    ]
    for p in candidates:
        if p.exists():
            return str(p)
    return None


def _find_unicode_font_bold() -> str | None:
    """Return path to the bold variant of a system TrueType font."""
    from pathlib import Path
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ]
    for p in candidates:
        if p.exists():
            return str(p)
    return None


def _compute_col_widths(cols: list[str], page_w: float) -> list[float]:
    """Return column widths proportional to weight, fitting page_w."""
    weights = [_COL_WEIGHT.get(k, _DEFAULT_WEIGHT) for k in cols]
    total = sum(weights)
    return [page_w * w / total for w in weights]


def _measure_cell_height(pdf, text: str, col_w: float, line_h: float) -> float:
    """Use fpdf2 dry_run to compute the exact height a multi_cell would need."""
    result = pdf.multi_cell(
        col_w - 1, line_h, text,
        border=0, align="L",
        dry_run=True, output="LINES",
    )
    n_lines = max(1, len(result))
    return n_lines * line_h + 1.0


def build_pdf(
    p1: list[dict] | None = None,
    p2: list[dict] | None = None,
    p3: list[dict] | None = None,
    verdict: dict | None = None,
    session_name: str = "",
) -> bytes:
    """Build a PDF governance report and return as bytes.
    Uses only approved phase data and verdict; feedback history and agent_message are not included.
    """
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=False)

    font_family = "Helvetica"
    font_path = _find_unicode_font()
    font_path_bold = _find_unicode_font_bold()
    if font_path:
        font_family = "ReportFont"
        pdf.add_font(font_family, "", font_path, uni=True)
        if font_path_bold:
            pdf.add_font(font_family, "B", font_path_bold, uni=True)
        else:
            pdf.add_font(font_family, "B", font_path, uni=True)

    bottom_margin = 15
    line_h = 3.5
    header_h = 7
    font_size_body = 6.5
    font_size_header = 7

    # ── Title page ────────────────────────────────────────────────────
    pdf.add_page()

    pdf.set_font(font_family, "B", 18)
    pdf.cell(0, 12, "Risk Governance Report", new_x="LMARGIN", new_y="NEXT", align="C")
    if session_name:
        pdf.set_font(font_family, "", 11)
        pdf.cell(0, 7, session_name, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font(font_family, "", 9)
    pdf.cell(
        0, 6,
        f"Generated: {datetime.now().strftime('%B %d, %Y %H:%M')}",
        new_x="LMARGIN", new_y="NEXT", align="C",
    )
    pdf.ln(6)

    # ── Table renderer ────────────────────────────────────────────────
    def _draw_header(cols, headers, col_widths):
        pdf.set_font(font_family, "B", font_size_header)
        pdf.set_fill_color(26, 26, 46)
        pdf.set_text_color(255, 255, 255)
        for j, h in enumerate(headers):
            pdf.cell(col_widths[j], header_h, h, border=1, fill=True, align="C")
        pdf.ln()
        pdf.set_text_color(0, 0, 0)

    def _table(title: str, data: list[dict], col_map: dict):
        if not data:
            return
        cols = [k for k in col_map if k == "Comment" or any(k in row for row in data)]
        headers = [col_map[k] for k in cols]
        page_w = pdf.w - pdf.l_margin - pdf.r_margin
        col_widths = _compute_col_widths(cols, page_w)

        pdf.set_font(font_family, "B", 13)
        pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

        _draw_header(cols, headers, col_widths)

        pdf.set_font(font_family, "", font_size_body)

        for i, row in enumerate(data):
            cell_texts = [_cell_value(row, key) for key in cols]

            row_h = 0
            for j, val in enumerate(cell_texts):
                h = _measure_cell_height(pdf, val, col_widths[j], line_h)
                row_h = max(row_h, h)
            row_h = max(5, row_h)

            if pdf.get_y() + row_h > pdf.h - bottom_margin:
                pdf.add_page()
                _draw_header(cols, headers, col_widths)
                pdf.set_font(font_family, "", font_size_body)

            y_start = pdf.get_y()
            fill_even = i % 2 == 0

            for j, val in enumerate(cell_texts):
                x = pdf.l_margin + sum(col_widths[:j])
                w = col_widths[j]

                if fill_even:
                    pdf.set_fill_color(245, 245, 250)
                    pdf.rect(x, y_start, w, row_h, style="DF")
                else:
                    pdf.set_fill_color(255, 255, 255)
                    pdf.rect(x, y_start, w, row_h, style="D")

                pdf.set_xy(x + 0.5, y_start + 0.5)
                pdf.multi_cell(w - 1, line_h, val, border=0, align="L")

            pdf.set_xy(pdf.l_margin, y_start + row_h)

        pdf.ln(8)

    # ── Phase tables ──────────────────────────────────────────────────
    if p1:
        _table("Phase 1 — Risk Register", p1, PHASE_1_COLS)
    if p2:
        pdf.add_page()
        _table("Phase 2 — Risk-Control Matrix", p2, PHASE_2_COLS)
    if p3:
        pdf.add_page()
        _table("Phase 3 — Governance Map", p3, PHASE_3_COLS)

    # ── Verdict ───────────────────────────────────────────────────────
    if verdict:
        pdf.add_page()
        pdf.set_font(font_family, "B", 14)
        pdf.cell(0, 10, "Governance Verdict", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

        status = verdict.get("overall_status", "Unknown")
        color_map = {"Ready": (46, 125, 50), "Partially Ready": (245, 124, 0), "Not Ready": (198, 40, 40)}
        r, g, b = color_map.get(status, (0, 0, 0))
        pdf.set_text_color(r, g, b)
        pdf.set_font(font_family, "B", 12)
        pdf.cell(0, 8, f"Overall Status: {status}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)

        verdict_text = verdict.get("overall_verdict", "")
        if verdict_text:
            pdf.set_font(font_family, "", 10)
            pdf.multi_cell(0, 5, verdict_text)
            pdf.ln(4)

        risks = verdict.get("risks", [])
        if risks:
            verdict_cols = {
                "risk_id": "Risk ID",
                "risk_statement": "Risk",
                "status": "Coverage",
                "reason": "Reason",
            }
            _table("Per-Risk Coverage", risks, verdict_cols)

        recs = verdict.get("recommendations", [])
        if recs:
            if pdf.get_y() > pdf.h - 60:
                pdf.add_page()
            pdf.set_font(font_family, "B", 12)
            pdf.cell(0, 8, "Recommendations", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
            pdf.set_font(font_family, "", 9)
            for idx, rec in enumerate(recs, 1):
                if pdf.get_y() > pdf.h - 20:
                    pdf.add_page()
                    pdf.set_font(font_family, "", 9)
                pdf.multi_cell(0, 5, f"{idx}. {rec}")
                pdf.ln(1)

    buf = io.BytesIO()
    pdf.output(buf)
    return buf.getvalue()
