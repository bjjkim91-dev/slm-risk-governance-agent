"""
Rebuild the input documents for training sets from the training records.

The handover package shipped the expected outputs and the training JSONL for five document
sets but the input files for only two of them (Set 2 and Set 3). Every training record
embeds the full text of its inputs, so this script writes the missing sets back out in the
layout prepare_training_data.py expects:

    data/test_docs/Set N/<regulation>.pdf        the regulation text, one page per [Page N] block
    data/test_docs/Set N/Control Inventory.xlsx  Control ID + Control Description rows
    data/test_docs/Set N/Process narratives.docx the process narrative, one paragraph per line

Regulations are public-domain US federal text; control inventories and narratives are
synthetic. Files are written with empty author metadata.

Usage:  python training_reference/rebuild_sample_sets.py [--sets 1 4 5]
"""
import argparse
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TRAINING = ROOT / "training_reference" / "training_data"
TEST_DOCS = ROOT / "data" / "test_docs"
SETS = ["Set 1", "Set 2", "Set 3", "Set 4", "Set 5"]


def _user_turn(record: dict) -> str:
    return next(m["content"] for m in record["conversations"] if m["role"] == "user")


def _between(text: str, start: str, end: str) -> str:
    a = text.index(start) + len(start)
    b = text.index(end, a)
    return text[a:b].strip("\n")


def _records(phase: str) -> list:
    with open(TRAINING / f"{phase}_train.jsonl", encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def write_regulation(set_dir: Path, set_idx: int, p1_records: list) -> Path:
    import pymupdf

    expected = json.load(open(ROOT / "training_reference" / SETS[set_idx] / "phase_1_expected.json", encoding="utf-8"))
    filename = expected[0]["source_document"]
    body = _between(_user_turn(p1_records[set_idx]), "--- DOCUMENT TEXT ---", "--- INSTRUCTION ---")
    pages = [p.strip("\n") for p in re.split(r"\[Page \d+\]\n?", body) if p.strip()]
    doc = pymupdf.open()
    for page_text in pages:
        page_text = page_text.replace("—", "-").replace("–", "-")  # the base font has no dashes
        page = doc.new_page(width=612, height=792)
        page.insert_textbox(pymupdf.Rect(54, 54, 558, 760), page_text, fontsize=8.5, fontname="helv")
    doc.set_metadata({})
    out = set_dir / filename
    doc.save(out)
    return out


def write_controls(set_dir: Path, set_idx: int, p2_records: list) -> Path:
    from openpyxl import Workbook

    body = _between(_user_turn(p2_records[set_idx]), "--- CONTROL INVENTORY ---", "--- INSTRUCTION ---")
    rows = re.findall(r"\[Row \d+\] Control ID: (\S+) \| Control Description: (.+)", body)
    wb = Workbook()
    ws = wb.active
    ws.title = "Controls"
    ws.append(["Control ID", "Control Description"])
    for cid, desc in rows:
        ws.append([cid, desc.strip()])
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    out = set_dir / "Control Inventory.xlsx"
    wb.save(out)
    return out


def write_processes(set_dir: Path, set_idx: int, p3_records: list) -> Path:
    from docx import Document

    body = _between(_user_turn(p3_records[set_idx]), "--- PROCESS DOCUMENTS ---", "--- INSTRUCTION ---")
    doc = Document()
    for line in body.splitlines():
        if line.strip():
            doc.add_paragraph(line.strip())
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    out = set_dir / "Process narratives.docx"
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sets", nargs="*", type=int, default=[1, 4, 5])
    args = ap.parse_args()
    p1, p2, p3 = _records("p1"), _records("p2"), _records("p3")
    for n in args.sets:
        set_dir = TEST_DOCS / f"Set {n}"
        set_dir.mkdir(parents=True, exist_ok=True)
        for out in (write_regulation(set_dir, n - 1, p1), write_controls(set_dir, n - 1, p2), write_processes(set_dir, n - 1, p3)):
            print(f"wrote {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
